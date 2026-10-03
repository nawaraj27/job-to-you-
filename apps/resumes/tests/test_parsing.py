"""
Deterministic reading (SRS §5.5, P4).

Real files, real readers. A parser tested only against strings it was written
for passes forever and fails on the first PDF a user uploads, so the format
tests go through pdfminer and python-docx rather than around them.
"""

from __future__ import annotations

from django.test import SimpleTestCase

from apps.files.constants import DOCX, PDF, PNG
from apps.resumes.parsing import ExtractionFailed, UnsupportedFormat, extract_text, parse
from apps.resumes.parsing.extract import NO_EXPERIENCE_WARNING, NO_TEXT_WARNING
from apps.resumes.parsing.text import page_count

from .factories import SAMPLE_CV, docx_bytes, pdf_bytes


class ExtractionTests(SimpleTestCase):
    def test_a_pdf_text_layer_is_read(self):
        text = extract_text(pdf_bytes("Priya Sharma\nBackend Engineer"), mime=PDF)

        self.assertIn("Priya Sharma", text)
        self.assertIn("Backend Engineer", text)

    def test_a_docx_is_read_including_its_tables(self):
        import io

        import docx

        document = docx.Document()
        document.add_paragraph("Priya Sharma")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Backend Engineer"
        table.rows[0].cells[1].text = "2023 - Present"
        buffer = io.BytesIO()
        document.save(buffer)

        text = extract_text(buffer.getvalue(), mime=DOCX)

        # A third of real CVs keep their dates in a table. A reader that skips
        # them reports "no experience" on a document that is only experience.
        self.assertIn("Backend Engineer", text)
        self.assertIn("2023 - Present", text)

    def test_a_corrupt_pdf_fails_cleanly(self):
        with self.assertRaises(ExtractionFailed):
            extract_text(b"%PDF-1.4 this is not a pdf", mime=PDF)

    def test_a_corrupt_docx_fails_cleanly(self):
        with self.assertRaises(ExtractionFailed):
            extract_text(b"PK\x03\x04 not really a docx", mime=DOCX)

    def test_an_unreadable_format_is_not_an_error_in_the_file(self):
        with self.assertRaises(UnsupportedFormat):
            extract_text(b"\x89PNG", mime=PNG)

    def test_page_count_is_reported_for_pdfs_and_skipped_elsewhere(self):
        self.assertEqual(page_count(pdf_bytes(), mime=PDF), 1)
        self.assertIsNone(page_count(docx_bytes(), mime=DOCX))

    def test_extraction_is_deterministic(self):
        data = pdf_bytes("Priya Sharma\nPython, Django")

        self.assertEqual(extract_text(data, mime=PDF), extract_text(data, mime=PDF))


class StructureTests(SimpleTestCase):
    def setUp(self):
        self.parsed = parse(SAMPLE_CV)

    def test_contact_details_are_matched_not_guessed(self):
        self.assertEqual(self.parsed.emails, ("priya.sharma@example.com",))
        self.assertIn("9812345678", self.parsed.phones[0])
        self.assertTrue(self.parsed.links)

    def test_a_line_of_profile_links_is_not_read_as_a_job_title(self):
        """A real CV's header, and what it did to a real profile.

        "linkedin.com/in/priya | github.com/priyash" is five words, is not a
        name, and contains no scheme — so it passed every filter and became the
        candidate's headline. Their own dashboard then greeted them with a line
        of URLs where their job title belongs.
        """
        parsed = parse(
            "Priya Sharma\n"
            "linkedin.com/in/priyasharma-be | github.com/priyash-dev\n"
            "Senior Backend Engineer\n"
        )

        self.assertEqual(parsed.headline, "Senior Backend Engineer")

    def test_the_candidates_own_name_is_read_off_the_cv(self):
        """The import took its email from the file and its name from the account.

        The pass that looks for a headline already identifies the name line in
        order to skip it, and threw the answer away — so a candidate whose
        account name differs from the one they wrote on their CV got an
        imported document carrying both.
        """
        parsed = parse("Bibek Shrestha\nSenior Data Engineer\nbibek@example.com\n")

        self.assertEqual(parsed.name, "Bibek Shrestha")
        self.assertEqual(parsed.headline, "Senior Data Engineer")
        self.assertEqual(parsed.as_structured()["name"], "Bibek Shrestha")

    def test_a_name_written_as_a_form_field_is_still_the_name(self):
        # A form-style CV writes every header field as "Label - value", and the
        # guard that stops "Name - sunita gurung" becoming a job title used to
        # stop it becoming a name as well: a whole class of first CV imported
        # with no name on it at all.
        parsed = parse("CURRICULUM VITAE\n\nName - sunita gurung\nemail s@example.com\n")

        self.assertEqual(parsed.name, "sunita gurung")

    def test_a_form_field_that_is_not_a_name_is_not_read_as_one(self):
        parsed = parse("CURRICULUM VITAE\n\nAddress - Lalitpur Nepal\nemail s@example.com\n")

        self.assertEqual(parsed.name, "")

    def test_a_cv_with_no_name_line_at_all_reports_none(self):
        parsed = parse("CURRICULUM VITAE\n\nemail s@example.com\n+977 9841 220 118\n")

        self.assertEqual(parsed.name, "")

    def test_a_form_label_is_not_read_as_a_job_title(self):
        # "Name - sunita gurung" is two to ten words, is not a name, and
        # objected to nothing — so it became the candidate's job title.
        parsed = parse(
            "CURRICULUM VITAE\n\nName - sunita gurung\naddress - kathmandu\n"
            "email sunitagurung123@example.com\n"
        )

        self.assertEqual(parsed.headline, "")

    def test_a_sentence_fragment_is_not_read_as_a_job_title(self):
        # A two-column CV extracts out of order, and the pieces start
        # mid-sentence. One of them became a designer's job title.
        parsed = parse(
            "A N J A L I   K A R K I\nhello@example.design\n"
            "slightly annoyed. Six years, mostly in logistics and\n"
        )

        self.assertEqual(parsed.headline, "")

    def test_a_short_strapline_is_accepted_as_a_job_title(self):
        parsed = parse("Anjali Karki\nProduct Designer\nhello@example.design\n")

        self.assertEqual(parsed.headline, "Product Designer")

    def test_a_year_is_not_mistaken_for_a_phone_number(self):
        parsed = parse("Worked 2018 - 2022 at Tribhuvan University")

        self.assertEqual(parsed.phones, ())

    def test_sections_are_split_on_headings(self):
        self.assertIn("experience", self.parsed.sections)
        self.assertIn("skills", self.parsed.sections)
        self.assertIn("education", self.parsed.sections)

    def test_a_heading_word_inside_a_sentence_does_not_start_a_section(self):
        parsed = parse("Summary\nExperience with distributed systems and payments.\n")

        self.assertIn("distributed systems", parsed.sections["summary"])
        self.assertNotIn("experience", parsed.sections)

    def test_skills_are_terms_not_resolved_skills(self):
        # The taxonomy decides what "Python" is. A parser that resolved skills
        # itself would be a second, worse resolver.
        self.assertIn("Python", self.parsed.skill_terms)
        self.assertIn("PostgreSQL", self.parsed.skill_terms)

    def test_roles_are_grouped_under_their_dated_line(self):
        first = self.parsed.experience[0]

        self.assertEqual(first.title, "Backend Engineer")
        self.assertEqual(first.org, "Nexa Payments")
        self.assertTrue(first.is_current)
        self.assertEqual(first.end, "")
        self.assertEqual(len(first.bullets), 2)

    def test_a_finished_role_keeps_its_end_date(self):
        second = self.parsed.experience[1]

        self.assertEqual(second.org, "Himal Software")
        self.assertFalse(second.is_current)
        self.assertIn("2023", second.end)

    def test_education_is_read(self):
        self.assertEqual(len(self.parsed.education), 1)
        self.assertIn("Tribhuvan", self.parsed.education[0].institution)

    def test_a_missing_section_is_absent_and_warned_about(self):
        parsed = parse("Priya Sharma\npriya@example.com\n")

        self.assertEqual(parsed.skill_terms, ())
        self.assertTrue(any("skills" in warning for warning in parsed.warnings))

    def test_a_scanned_file_is_told_it_is_a_picture(self):
        """The single most useful thing this product can say about a file.

        A scan or a phone photograph extracts to nothing. It used to be
        reported as "No experience section found" and "No skills section
        found" — which reads as a criticism of the candidate's CV rather than
        a fact about their file, and buries the one thing that matters: no
        applicant tracking system will read it either.
        """
        parsed = parse("Page 1 of 2\n")

        self.assertEqual(parsed.warnings, (NO_TEXT_WARNING,))
        self.assertEqual(parsed.experience, ())

    def test_a_short_document_with_contact_details_is_not_a_picture(self):
        # Every real CV carries an email, a phone number or a link, and none of
        # the three survives being photographed. Telling somebody holding a
        # real document that it is a photograph is the worse error.
        parsed = parse("Priya Sharma\npriya@example.com\n")

        self.assertNotIn(NO_TEXT_WARNING, parsed.warnings)

    def test_a_bare_work_heading_still_finds_the_experience(self):
        # "WORK" is what plenty of CVs write, and a heading this vocabulary
        # does not know costs the candidate their entire history: no roles
        # extracted, and a report telling them they have no experience.
        parsed = parse(
            "Sunita Gurung\nsunita@example.com\n\n"
            "WORK\nOfficer, Nabil Bank  2019 - 2023\n- Opened accounts.\n\n"
            "EDUCATION\nBBA, Kathmandu University  2015 - 2019\n"
        )

        self.assertIn("experience", parsed.sections)
        self.assertNotIn(NO_EXPERIENCE_WARNING, parsed.warnings)

    def test_confidence_is_higher_for_matched_patterns_than_for_position(self):
        self.assertGreater(self.parsed.confidence["contact"], self.parsed.confidence["experience"])

    def test_empty_input_parses_to_nothing_without_raising(self):
        parsed = parse("")

        self.assertEqual(parsed.skill_terms, ())
        self.assertEqual(parsed.experience, ())

    def test_the_structured_shape_is_json_serialisable(self):
        import json

        json.dumps(self.parsed.as_structured())

    def test_parsing_is_deterministic(self):
        self.assertEqual(parse(SAMPLE_CV).as_structured(), parse(SAMPLE_CV).as_structured())


class SectionBoundaryTests(SimpleTestCase):
    """A heading this reader does not know must still end the section above it.

    Every case here produced content the candidate never claimed. That is the
    worst class of error the reader can make: a skill chip, an extra degree or
    a second job appears on their profile, attributed to them by the platform,
    and nothing in the product knows it is wrong.
    """

    def test_an_unknown_heading_ends_the_section_above_it(self):
        parsed = parse(
            "Kiran Shrestha\nkiran@example.com\n\n"
            "SKILLS\nMS Word, MS Excel\n\n"
            "REFERENCE\nAvailable upon request\n"
        )

        self.assertEqual(parsed.skill_terms, ("MS Word", "MS Excel"))

    def test_a_publication_list_does_not_become_education(self):
        parsed = parse(
            "Dr Bimala Adhikari\nbimala@example.edu.np\n\n"
            "EDUCATION\nPhD Glaciology, University of Bergen, 2014 - 2018\n\n"
            "PUBLICATIONS\nAdhikari B. (2024) Debris-covered glacier retreat.\n"
        )

        self.assertEqual([row.qualification for row in parsed.education], ["PhD Glaciology"])

    def test_a_skill_on_its_own_line_is_not_read_as_a_heading(self):
        # The guard above only fires after a blank line. Without that, a CV
        # that lists one skill per line loses every skill under the first
        # single-word one.
        parsed = parse("Sneha\nsneha@example.com\n\nTOOLS\nSQL\nPower BI\nPython\n")

        self.assertEqual(parsed.skill_terms, ("SQL", "Power BI", "Python"))

    def test_an_institution_line_is_not_read_as_a_heading(self):
        parsed = parse(
            "Deepak Poudel\ndeepak@example.com\n\n"
            "EDUCATION\nMaster of Education\nTribhuvan University\n2012 - 2014\n"
        )

        self.assertEqual(len(parsed.education), 1)
        self.assertEqual(parsed.education[0].institution, "Tribhuvan University")


class ExperienceShapeTests(SimpleTestCase):
    def test_dates_written_under_the_role_still_find_the_role(self):
        # The most common Word layout there is. Reading the dated line as the
        # role gave the candidate a job with no title and told them they had
        # no work experience.
        parsed = parse(
            "Sunita Gurung\nsunita@example.com\n\n"
            "PROFESSIONAL EXPERIENCE\n"
            "Staff Nurse — Patan Hospital\n2017 — 2020\n"
        )

        self.assertEqual(len(parsed.experience), 1)
        self.assertEqual(parsed.experience[0].title, "Staff Nurse")
        self.assertEqual(parsed.experience[0].org, "Patan Hospital")
        self.assertEqual(parsed.experience[0].start, "2017")

    def test_dates_written_above_the_role_still_find_the_role(self):
        parsed = parse(
            "Manisha\nmanisha@example.com\n\n"
            "EXPERIENCE\n2022 - Present\nProduct Designer, Foodmandu\n"
        )

        self.assertEqual(len(parsed.experience), 1)
        self.assertEqual(parsed.experience[0].title, "Product Designer")
        self.assertTrue(parsed.experience[0].is_current)

    def test_internships_are_experience(self):
        parsed = parse(
            "Anisha Thapa\nanisha@example.com\n\n"
            "INTERNSHIPS\nFinance Intern, NIC Asia Bank — Jun 2025 to Sep 2025\n"
        )

        self.assertEqual(len(parsed.experience), 1)
        self.assertNotIn(NO_EXPERIENCE_WARNING, parsed.warnings)

    def test_a_cv_with_no_employment_reports_none(self):
        parsed = parse(
            "Sabin Bhandari\nsabin@example.com\n\n"
            "EDUCATION\nBSc Computing, Islington College, 2022 - 2026\n\n"
            "PROJECTS\nBus Tracker — a live map, built with Flutter.\n\n"
            "SKILLS\nPython, Java\n"
        )

        self.assertEqual(parsed.experience, ())


class ContactShapeTests(SimpleTestCase):
    def test_an_orcid_is_not_a_telephone_number(self):
        # Sixteen digits with separators matches any phone pattern loose enough
        # for the formats this launches into, so an academic's ORCID was shown
        # to employers as their phone number.
        parsed = parse("Dr Bimala Adhikari\nbimala@example.edu.np\norcid.org/0000-0002-1825-0097\n")

        self.assertEqual(parsed.phones, ())
        self.assertEqual(parsed.links, ("orcid.org/0000-0002-1825-0097",))

    def test_a_real_phone_number_beside_a_link_is_still_found(self):
        parsed = parse("Priya\npriya@example.com | +977 9841 220 118 | github.com/priya/cv\n")

        self.assertEqual(parsed.phones, ("+977 9841 220 118",))


class EmptyFileTests(SimpleTestCase):
    def test_an_empty_document_is_not_called_a_photograph(self):
        from apps.resumes.parsing.extract import EMPTY_WARNING

        parsed = parse("Untitled\n")

        self.assertEqual(parsed.warnings, (EMPTY_WARNING,))

    def test_page_furniture_alone_is_still_a_photograph(self):
        parsed = parse("Page 1 of 2\n")

        self.assertEqual(parsed.warnings, (NO_TEXT_WARNING,))


class LayoutTests(SimpleTestCase):
    """A PDF has no reading order. Guessing one wrongly loses whole sections."""

    def test_a_sidebar_label_is_read_before_the_column_beside_it(self):
        import io

        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas

        buffer = io.BytesIO()
        _width, height = A4
        pdf = canvas.Canvas(buffer, pagesize=A4)
        pdf.setFont("Helvetica", 9)
        pdf.drawString(180, height - 60, "Anjali Karki")
        pdf.drawString(180, height - 73, "hello@example.design")
        pdf.drawString(54, height - 100, "WORK")
        pdf.drawString(180, height - 100, "Lead Product Designer")
        pdf.drawRightString(540, height - 100, "2022 - now")
        pdf.drawString(180, height - 113, "Pathao Nepal")
        pdf.drawString(54, height - 140, "TOOLS")
        pdf.drawString(180, height - 140, "Figma, Framer, Maze")
        pdf.showPage()
        pdf.save()

        parsed = parse(extract_text(buffer.getvalue(), mime=PDF))

        # Flat extraction collected every label at the bottom of the file,
        # detached from the content it labelled: no headings were found, and a
        # designer's CV read as no skills, no roles and no education at all.
        self.assertEqual(parsed.skill_terms, ("Figma", "Framer", "Maze"))
        self.assertEqual(len(parsed.experience), 1)
        self.assertEqual(parsed.experience[0].title, "Lead Product Designer")
        self.assertEqual(parsed.experience[0].org, "Pathao Nepal")
        self.assertTrue(parsed.experience[0].is_current)

    def test_a_table_row_keeps_its_cells_apart(self):
        import io

        import docx

        document = docx.Document()
        document.add_paragraph("Laxmi Bhattarai")
        document.add_paragraph("WORK EXPERIENCE")
        table = document.add_table(rows=1, cols=3)
        cells = table.rows[0].cells
        cells[0].text = "Primary School Teacher"
        cells[1].text = "Rato Bangala School"
        cells[2].text = "2018 - Present"
        buffer = io.BytesIO()
        document.save(buffer)

        parsed = parse(extract_text(buffer.getvalue(), mime=DOCX))

        # Joined with whitespace, the three cells became one job title and the
        # role reached the profile with no employer on it.
        self.assertEqual(len(parsed.experience), 1)
        self.assertEqual(parsed.experience[0].title, "Primary School Teacher")
        self.assertEqual(parsed.experience[0].org, "Rato Bangala School")

    def test_a_table_is_read_where_the_document_puts_it(self):
        import io

        import docx

        document = docx.Document()
        document.add_paragraph("Sita Rai")
        document.add_paragraph("EXPERIENCE")
        table = document.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Marketing Manager, Daraz"
        table.rows[0].cells[1].text = "2021 - Present"
        document.add_paragraph("SKILLS")
        document.add_paragraph("SEO, Google Ads")
        buffer = io.BytesIO()
        document.save(buffer)

        parsed = parse(extract_text(buffer.getvalue(), mime=DOCX))

        # Reading the paragraphs first and appending the tables afterwards put
        # the employment history below the final heading, so a Word CV with its
        # roles in a table had its whole career read as a list of skills.
        self.assertEqual(parsed.skill_terms, ("SEO", "Google Ads"))
        self.assertEqual(len(parsed.experience), 1)


class BulletedListTests(SimpleTestCase):
    def test_a_bulleted_skills_list_with_gaps_keeps_every_skill(self):
        # How the Word reader delivers a bulleted list: a blank line between
        # items. "Docker" and "GitHub Actions" look like headings on their own.
        parsed = parse(
            "Maya Karki\nmaya@example.com\n\n"
            "EXPERIENCE\nBackend Developer, Phewa Stays - Jan 2024 - Present\n\n"
            "• Built booking APIs.\n"
            "SKILLS\n\n• Python\n\n• Django\n\n• Docker\n\n• Git\n\n• GitHub Actions\n"
            "EDUCATION\nBSc CSIT, Tribhuvan University, 2018 - 2022\n"
        )

        self.assertEqual(
            parsed.skill_terms, ("Python", "Django", "Docker", "Git", "GitHub Actions")
        )
        self.assertEqual(len(parsed.education), 1)


class EvidenceSectionTests(SimpleTestCase):
    """Projects, certifications, awards and publications.

    Read for a long time and carried nowhere. These are the sections a CV
    built out of anything other than employment depends on, which is every
    student, graduate, career changer and academic on the platform.
    """

    def test_projects_keep_their_descriptions(self):
        parsed = parse(
            "Sabin\nsabin@example.com\n\n"
            "PROJECTS\n"
            "Bus Tracker - a live map of valley routes, built with Flutter.\n"
            "Notes API\n"
            "- A REST API with token authentication.\n"
        )

        self.assertEqual([row.name for row in parsed.projects], ["Bus Tracker", "Notes API"])
        self.assertEqual(
            parsed.projects[0].bullets, ("a live map of valley routes, built with Flutter.",)
        )
        self.assertEqual(parsed.projects[1].bullets, ("A REST API with token authentication.",))

    def test_a_bulleted_list_of_projects_keeps_every_project(self):
        # No name lines, only bullets: each bullet used to need a project above
        # it to attach to, so the whole section was dropped.
        parsed = parse(
            "Maya\nmaya@example.com\n\n"
            "PROJECTS\n"
            "- Room availability cache with Redis; reduced database reads by 70%.\n"
            "- Bus Tracker - a live map of valley routes.\n"
        )

        self.assertEqual(
            [row.name for row in parsed.projects],
            ["Room availability cache with Redis; reduced database reads by 70%.", "Bus Tracker"],
        )
        self.assertEqual(parsed.projects[1].bullets, ("a live map of valley routes.",))

    def test_certifications_and_awards_are_listed(self):
        parsed = parse(
            "Anisha\nanisha@example.com\n\n"
            "CERTIFICATIONS\nCFA Level I\nBloomberg Market Concepts\n\n"
            "AWARDS\nDean's list, 2025\n"
        )

        self.assertEqual(parsed.certifications, ("CFA Level I", "Bloomberg Market Concepts"))
        self.assertEqual(parsed.awards, ("Dean's list, 2025",))

    def test_a_numbered_publication_list_keeps_each_citation_whole(self):
        # Academic CVs number their publications and every citation wraps. The
        # continuation of a citation starts with a capital, so the bullet rule
        # read it as a new item: each publication split in two and the next
        # one's number glued itself onto the stray half.
        parsed = parse(
            "Dr Prakash\nprakash@example.edu.np\n\n"
            "SELECTED PUBLICATIONS\n"
            "1. Adhikari, P. (2011). Snowmelt contribution to dry-season discharge.\n"
            "Journal of Himalayan Hydrology, 11(2), 107-119.\n"
            "2. Adhikari, P. (2014). Debris cover and glacier retreat.\n"
            "Cryosphere Letters, 6(1), 22-31.\n"
        )

        self.assertEqual(len(parsed.publications), 2)
        self.assertTrue(parsed.publications[0].startswith("Adhikari, P. (2011)"))
        self.assertIn("Journal of Himalayan Hydrology", parsed.publications[0])
        self.assertTrue(parsed.publications[1].startswith("Adhikari, P. (2014)"))

    def test_grants_are_read_as_awards(self):
        parsed = parse("Dr Prakash\np@example.edu.np\n\nGRANTS\nNepal Academy of Science, 2021\n")

        self.assertEqual(parsed.awards, ("Nepal Academy of Science, 2021",))

    def test_a_cv_without_them_reports_none(self):
        parsed = parse("Priya\npriya@example.com\n\nSKILLS\nPython, Django\n")

        self.assertEqual(parsed.projects, ())
        self.assertEqual(parsed.certifications, ())
        self.assertEqual(parsed.awards, ())
        self.assertEqual(parsed.publications, ())

    def test_a_word_list_paragraph_keeps_its_bullet(self):
        """Word draws the bullet from a style and puts none in the text.

        Every reader downstream that asks "is this a bullet?" got no, for a
        document that is nothing but bullets — and a student's two project
        bullets were read as two more projects.
        """
        import io

        import docx

        document = docx.Document()
        document.add_paragraph("Sabin")
        document.add_paragraph("sabin@example.com")
        document.add_paragraph("PROJECTS")
        document.add_paragraph("Notes API")
        document.add_paragraph("A REST API with token authentication.", style="List Bullet")
        document.add_paragraph("Deployed on a free tier.", style="List Bullet")
        buffer = io.BytesIO()
        document.save(buffer)

        parsed = parse(extract_text(buffer.getvalue(), mime=DOCX))

        self.assertEqual([row.name for row in parsed.projects], ["Notes API"])
        self.assertEqual(len(parsed.projects[0].bullets), 2)


class NamesInEveryScriptTests(SimpleTestCase):
    """ "Capitalised" is a property of cased scripts only.

    Devanagari, Arabic, Thai, Chinese, Japanese and Korean have no case at all,
    so `isupper()` is False for every character in them — and a name test
    written as "each word is capitalised" rejects every name written in any of
    them. That is not a niche here: this reader's own vocabulary is tuned for
    Nepal and launches into that market first, where नवराज गौतम is an ordinary
    way to write a name. Every such candidate imported a CV with no name on it.
    """

    HEADER = "Senior Backend Engineer\np@example.com | 9841 220 118\n"
    BODY = "\nEXPERIENCE\nEngineer, Khalti - 2021 to Present\n\nSKILLS\nPython, Django\n"

    def _parse(self, name: str):
        return parse(f"{name}\n{self.HEADER}{self.BODY}")

    def test_a_name_is_read_in_every_script_a_candidate_might_use(self):
        for label, name in (
            ("Latin", "Priya Sharma"),
            ("Devanagari", "नवराज गौतम"),
            ("Arabic", "بريا شارما"),
            ("Thai", "สมชาย ใจดี"),
            ("Cyrillic", "Иван Петров"),
            ("Greek", "Γιώργος Παπάς"),
            ("Chinese", "张伟"),
            ("Japanese", "山田 太郎"),
            ("Korean", "김민준"),
        ):
            with self.subTest(script=label):
                parsed = self._parse(name)

                self.assertEqual(parsed.name, name)
                # And reading the name must not cost the headline underneath.
                self.assertEqual(parsed.headline, "Senior Backend Engineer")

    def test_a_lone_latin_word_is_still_not_a_name(self):
        # One word is normally a heading. The single-token allowance is for
        # scripts that write a whole name that way, and must not open the door
        # to "CURRICULUM VITAE" or "EXPERIENCE" becoming somebody's name.
        parsed = parse("CURRICULUM VITAE\n\nEXPERIENCE\nEngineer, Khalti - 2021 to Present\n")

        self.assertEqual(parsed.name, "")

    def test_a_job_title_is_still_not_read_as_a_name(self):
        parsed = parse("Marketing Manager\np@example.com\n")

        self.assertEqual(parsed.name, "")


class EmptyWordFileTests(SimpleTestCase):
    def test_an_empty_word_file_is_not_called_a_photograph(self):
        """A Word file cannot be a scan.

        Somebody who uploaded an empty .docx was told their file was "a picture
        of a CV rather than a document", and advised to export it from Word as
        a PDF instead of scanning it — advice about a machine they never
        touched, for a document they wrote in Word.
        """
        import io

        import docx

        from apps.resumes.parsing.extract import EMPTY_WARNING

        document = docx.Document()
        buffer = io.BytesIO()
        document.save(buffer)

        parsed = parse(extract_text(buffer.getvalue(), mime=DOCX), mime=DOCX)

        self.assertIn(EMPTY_WARNING, parsed.warnings)
        self.assertNotIn(NO_TEXT_WARNING, parsed.warnings)

    def test_a_scanned_pdf_is_still_called_a_photograph(self):
        parsed = parse("Page 1 of 2\n", mime=PDF)

        self.assertEqual(parsed.warnings, (NO_TEXT_WARNING,))


class ContactLocationTests(SimpleTestCase):
    """The place written beside the email or phone — read, never guessed."""

    def test_a_city_and_country_on_the_contact_line_is_read(self):
        parsed = parse(
            "Maya Karki\nmaya@example.com | +977 9800000000 | Pokhara, Nepal\n\nSKILLS\nPython\n"
        )
        self.assertEqual(parsed.location, "Pokhara, Nepal")
        self.assertEqual(parsed.as_structured()["contact"]["location"], "Pokhara, Nepal")

    def test_a_single_place_name_on_the_contact_line_is_read(self):
        parsed = parse(
            "Anjali\nhello@anjali.design · anjali.design/work · Pokhara\n\nSKILLS\nFigma\n"
        )
        self.assertEqual(parsed.location, "Pokhara")

    def test_a_comma_phrase_off_the_contact_line_is_not_a_place(self):
        parsed = parse(
            "Sita Rai\nsita@example.com | 9812345678\n\nEXPERIENCE\nEngineer, Fonepay - 2021 - Present\n"
        )
        self.assertEqual(parsed.location, "")

    def test_link_labels_and_street_numbers_are_not_places(self):
        for header in (
            "ram@example.com | 9812345678 | linkedin.com/in/ram | Portfolio",
            "ram@example.com | 9812345678 | Lakeside-6, Pokhara 33700",
            "ram@example.com | 9812345678 | Available for relocation",
        ):
            with self.subTest(header=header):
                self.assertEqual(parse(f"Ram Thapa\n{header}\n\nSKILLS\nPython\n").location, "")

    def test_no_contact_line_means_no_location(self):
        self.assertEqual(parse("Ram Thapa\nKathmandu, Nepal\n\nSKILLS\nPython\n").location, "")
